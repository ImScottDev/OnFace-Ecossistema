#!/usr/bin/env python3
from __future__ import annotations

import csv
import hashlib
import json
import re
import subprocess
import tempfile
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

SOURCE = Path("/Users/Scott/Downloads/auditoria-callface-zoho-crm2/reconstrucao-cliente-360/03_faturamento_mensal_por_cliente.xlsx")


def norm_doc(value):
    return re.sub(r"\D+", "", str(value or ""))


def dec(value):
    if value is None or value == "":
        return Decimal("0")
    return Decimal(str(value).replace(".", "").replace(",", ".") if "," in str(value) else str(value))


def run_psql(sql=None, file=None, capture=False):
    cmd = ["psql", "-U", "callface", "-d", "callface_crm", "-X", "-v", "ON_ERROR_STOP=1"]
    if capture:
        cmd += ["--csv", "-q", "-c", sql]
        p = subprocess.run(cmd, text=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    else:
        cmd += ["-f", str(file)]
        p = subprocess.run(cmd, text=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    if p.returncode:
        raise SystemExit(p.stderr)
    return p.stdout


def classify(row):
    family = str(row.get("familia_evidencia") or "").strip()
    nature = str(row.get("natureza_evidencia") or "").strip()
    business = str(row.get("business_line") or "").upper()
    if family and family != "UNKNOWN":
        return family, nature if nature and nature != "receita_emitida_sem_produto" else "MRR", family, "CONFIRMED"
    if "CALL FACE" in business or "PABX" in business or "CLOUD" in business:
        return "PABX", "MRR", "PABX", "PROBABLE"
    if "AI" in business:
        return "CALLFACE_AI", "MRR", "Callface AI", "PROBABLE"
    if "WHITE" in business:
        return "WHITE_LABEL", "MRR", "White Label", "PROBABLE"
    return "UNKNOWN", "UNKNOWN", None, "UNRESOLVED"


def main():
    org = next(csv.DictReader(run_psql("SELECT id FROM organizacoes ORDER BY created_at LIMIT 1", capture=True).splitlines()))["id"]
    companies_by_doc = {}
    for r in csv.DictReader(run_psql("SELECT id::text, regexp_replace(coalesce(cnpj,''),'\\D','','g') AS doc FROM empresas WHERE cnpj IS NOT NULL", capture=True).splitlines()):
        if not r["doc"]:
            continue
        companies_by_doc.setdefault(r["doc"], []).append(r["id"])

    wb = load_workbook(SOURCE, read_only=True, data_only=True)
    ws = wb["linhas_fontes"]
    it = ws.iter_rows(values_only=True)
    headers = [str(h or "") for h in next(it)]
    rows = []
    for raw in it:
        row = {headers[i]: raw[i] if i < len(raw) else None for i in range(len(headers))}
        doc = norm_doc(row.get("CNPJ"))
        month = str(row.get("mes") or "")
        if not month:
            continue
        competence = month + "-01" if len(month) == 7 else month[:10]
        value = dec(row.get("valor_faturado"))
        if value == 0:
            continue
        family, nature, product, confidence = classify(row)
        company_ids = companies_by_doc.get(doc, [])
        empresa_id = company_ids[0] if len(company_ids) == 1 else None
        source_hash = hashlib.sha256(json.dumps({
            "source": row.get("fonte"),
            "row": row.get("linha_origem"),
            "customer_key": row.get("customer_key"),
            "month": competence,
            "value": str(value),
            "external_id": row.get("external_id"),
            "family": family,
        }, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
        evidence = {
            "fonte": row.get("fonte"),
            "linha_origem": row.get("linha_origem"),
            "business_line": row.get("business_line"),
            "legal_entity": row.get("legal_entity"),
            "status_fatura": row.get("status_fatura"),
            "external_id": row.get("external_id"),
            "match": "cnpj_unico" if empresa_id else "sem_empresa_unica",
        }
        rows.append({
            "organizacao_id": org,
            "empresa_id": empresa_id or "",
            "customer_key": row.get("customer_key") or "",
            "cliente_nome": row.get("cliente") or row.get("customer_key") or "Cliente sem nome",
            "documento": doc,
            "competencia": competence,
            "valor": f"{value:.2f}",
            "quantidade_lancamentos": int(row.get("quantidade_notas") or 1),
            "service_family": family,
            "revenue_nature": nature,
            "canonical_product": product or "",
            "confidence": confidence,
            "source": "reconstrucao_cliente_360_linhas_fontes",
            "source_file": row.get("fonte") or SOURCE.name,
            "source_row": str(row.get("linha_origem") or ""),
            "evidence": json.dumps(evidence, ensure_ascii=False),
            "source_hash": source_hash,
        })
    wb.close()

    with tempfile.NamedTemporaryFile("w", newline="", suffix=".csv", delete=False) as fh:
        writer = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)
        csv_path = fh.name

    sql = f"""
BEGIN;
CREATE TEMP TABLE tmp_cliente_receita_mensal (LIKE cliente_receita_mensal INCLUDING DEFAULTS);
\\copy tmp_cliente_receita_mensal(organizacao_id,empresa_id,customer_key,cliente_nome,documento,competencia,valor,quantidade_lancamentos,service_family,revenue_nature,canonical_product,confidence,source,source_file,source_row,evidence,source_hash) FROM '{csv_path}' WITH CSV HEADER
INSERT INTO cliente_receita_mensal(organizacao_id,empresa_id,customer_key,cliente_nome,documento,competencia,valor,quantidade_lancamentos,service_family,revenue_nature,canonical_product,confidence,source,source_file,source_row,evidence,source_hash)
SELECT organizacao_id, NULLIF(empresa_id::text,'')::uuid, customer_key, cliente_nome, NULLIF(documento,''), competencia, valor, quantidade_lancamentos,
       service_family, revenue_nature, NULLIF(canonical_product,''), confidence, source, source_file, source_row, evidence, source_hash
FROM tmp_cliente_receita_mensal
ON CONFLICT (organizacao_id, source_hash) DO UPDATE SET
  empresa_id=EXCLUDED.empresa_id,
  cliente_nome=EXCLUDED.cliente_nome,
  documento=EXCLUDED.documento,
  valor=EXCLUDED.valor,
  quantidade_lancamentos=EXCLUDED.quantidade_lancamentos,
  service_family=EXCLUDED.service_family,
  revenue_nature=EXCLUDED.revenue_nature,
  canonical_product=EXCLUDED.canonical_product,
  confidence=EXCLUDED.confidence,
  evidence=EXCLUDED.evidence,
  updated_at=now();
COMMIT;
"""
    with tempfile.NamedTemporaryFile("w", suffix=".sql", delete=False) as fh:
        fh.write(sql)
        sql_path = fh.name
    run_psql(file=sql_path)
    print(json.dumps({"imported_rows": len(rows), "source": str(SOURCE)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
