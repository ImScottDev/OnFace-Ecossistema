#!/usr/bin/env python3
"""
Catálogo Oficial Stripe -> CRM2, passo 1 de 2: planilha -> snapshot JSON.

Por que dois passos em vez de um script TS lendo o .xlsx direto: o projeto não
tem nenhum parser de xlsx instalado (verificado no package.json), e a regra é
não instalar biblioteca sem necessidade real. O Python desta máquina já tem
openpyxl. O JSON gerado fica VERSIONADO no repo, o que dá três coisas de
graça: o import roda em qualquer máquina sem a planilha, dá `git diff` quando
o Kauã mudar um preço, e o snapshot é auditável (é catálogo comercial, não
dado de cliente — pode entrar no Git).

  python3 scripts/converter-catalogo-oficial.py <planilha.xlsx>
"""
import json, sys, unicodedata, re, collections
from pathlib import Path
import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
DESTINO = RAIZ / "app/src/lib/db/catalogo/catalogo-oficial.json"

FAMILIA = {"Callface AI": "CALLFACE_AI", "White Label": "WHITE_LABEL", "PABX": "PABX"}

# Categoria CRM -> slug estável. É a chave de agrupamento do builder, então
# precisa ser previsível: nunca derivada de acento/caixa em runtime.
CATEGORIA = {
    "Plano": "plano", "Add-on": "add_on", "Uso": "uso", "Setup": "setup",
    "Licença": "licenca", "Linha": "linha", "Ramal": "ramal", "Canais": "canais",
    "Hardware": "hardware", "Integração": "integracao", "Minutos": "minutos",
    "Serviço": "servico",
}

# Ordem de exibição por família no builder — do mais vendido/estruturante pro
# acessório. Não é alfabética de propósito: "Plano" tem que vir primeiro.
ORDEM_CATEGORIA = ["plano", "licenca", "setup", "add_on", "linha", "ramal", "canais",
                   "minutos", "uso", "hardware", "integracao", "servico"]


def slug(v):
    if not v:
        return ""
    s = unicodedata.normalize("NFKD", str(v)).encode("ascii", "ignore").decode()
    return re.sub(r"_+", "_", re.sub(r"[^a-z0-9]+", "_", s.lower())).strip("_")


# "Variação vinda do Zoho: NOME | COMPLEMENTO" — 189 dos 213 prices carregam o
# nome LEGADO exato na observação operacional. É um mapa legado->canônico
# confirmado pelo próprio time comercial, e vale mais que qualquer regra por
# palavra-chave: por isso vira `nomesLegado` no snapshot e alimenta a
# classificação do histórico com confiança CONFIRMED.
RE_ZOHO = re.compile(r"Varia[çc][ãa]o vinda do Zoho:\s*(.+)", re.IGNORECASE)


def nomes_legado(observacao):
    if not observacao:
        return []
    m = RE_ZOHO.search(str(observacao))
    if not m:
        return []
    bruto = m.group(1).strip()
    # O complemento depois do "|" descreve a condição comercial ("LOCACAO",
    # "TARIFA DESCRITA NA PROPOSTA"), não o produto. O nome legado real é o
    # primeiro pedaço; guardamos o texto inteiro também porque em contrato
    # antigo o nome às vezes foi gravado com o complemento junto.
    partes = [p.strip() for p in bruto.split("|") if p.strip()]
    saida = []
    for cand in ([partes[0]] if partes else []) + [bruto]:
        if cand and cand not in saida:
            saida.append(cand)
    return saida


def natureza_receita(tipo_receita, categoria_slug):
    """
    tipo_receita da planilha -> revenue_nature do CRM.

    A regra dura do Kauã (invariante do portfólio, ver servico-canonico.ts) é
    que SETUP, HARDWARE e USAGE nunca alimentam current_recurring_value. Por
    isso locação MENSAL de hardware é classificada HARDWARE, não MRR, mesmo
    sendo receita recorrente de fato — 34 dos 45 prices de PABX Hardware são
    locação mensal. Isso mantém a invariante intacta, mas significa que essa
    receita fica FORA do MRR do portfólio: é decisão comercial pendente,
    registrada em crm2-decisoes-pendentes.md, não um bug de classificação.
    """
    if tipo_receita == "setup":
        return "SETUP"
    if tipo_receita == "uso":
        return "USAGE"
    if categoria_slug == "hardware":
        return "HARDWARE"
    if tipo_receita == "recorrente":
        return "MRR"
    return "ONE_TIME"


def tipo_item(tipo_preco, tipo_receita):
    """tipo do item de proposta: recorrente | setup | avulso (enum já existente)."""
    if tipo_receita == "setup":
        return "setup"
    return "recorrente" if tipo_preco == "Recorrente" else "avulso"


def comportamento_quantidade(categoria_slug, produto_nome):
    """
    quantity_behavior (coluna que já existe em servicos_catalogo):
      LICENSE_SEATS -> quantidade = licenças (White Label)
      CAPACITY      -> quantidade = capacidade contratada (ramal/linha/canal)
      UNIT          -> quantidade = unidades
    """
    if categoria_slug in ("licenca",) or "White Label" in (produto_nome or ""):
        return "LICENSE_SEATS"
    if categoria_slug in ("ramal", "linha", "canais"):
        return "CAPACITY"
    return "UNIT"


def main():
    origem = Path(sys.argv[1] if len(sys.argv) > 1 else "")
    if not origem.exists():
        sys.exit(f"planilha não encontrada: {origem}")

    ws = openpyxl.load_workbook(origem, data_only=True)["Catalogo_Unico"]
    hdr = [c.value for c in ws[1]]
    linhas = [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True)
              if not all(v is None for v in r)]

    produtos = collections.OrderedDict()
    problemas, avisos = [], []

    for i, r in enumerate(linhas, start=2):
        pid, prid = r["STRIPE_PRODUCT_ID_COLAR_AQUI"], r["STRIPE_PRICE_ID_COLAR_AQUI"]
        if not pid or not str(pid).startswith("prod_"):
            problemas.append(f"linha {i}: STRIPE_PRODUCT_ID inválido ({pid!r})")
            continue
        if not prid or not str(prid).startswith("price_"):
            problemas.append(f"linha {i}: STRIPE_PRICE_ID inválido ({prid!r})")
            continue
        fam = FAMILIA.get(r["FAMILIA"])
        if not fam:
            problemas.append(f"linha {i}: FAMILIA desconhecida ({r['FAMILIA']!r})")
            continue
        cat = CATEGORIA.get(r["CATEGORIA_CRM"])
        if not cat:
            problemas.append(f"linha {i}: CATEGORIA_CRM desconhecida ({r['CATEGORIA_CRM']!r})")
            continue
        if r["VALOR_UNITARIO"] is None:
            problemas.append(f"linha {i}: VALOR_UNITARIO vazio ({prid})")
            continue

        p = produtos.setdefault(pid, {
            "stripeProductId": str(pid),
            "nome": r["NOME_PRODUTO_STRIPE"],
            "descricao": r["DESCRICAO_PRODUTO"],
            "descricaoExtrato": r["DESCRICAO_EXTRATO"],
            "etiquetaUnidade": r["ETIQUETA_UNIDADE"],
            "serviceFamily": fam,
            "categoria": cat,
            "categoriaLabel": r["CATEGORIA_CRM"],
            "skuProduto": slug(r["NOME_PRODUTO_STRIPE"]),
            "destinoPosVenda": r["METADATA_destino_pos_venda"],
            "observacaoOperacional": r["OBSERVACAO_OPERACIONAL"],
            "comoApareceNoCrm": r["COMO_APARECE_NO_CRM"],
            "quantityBehavior": comportamento_quantidade(cat, r["NOME_PRODUTO_STRIPE"]),
            "ordem": ORDEM_CATEGORIA.index(cat) * 100 if cat in ORDEM_CATEGORIA else 9900,
            "precos": [],
        })

        tipo_receita = r["METADATA_tipo_receita"]
        p["precos"].append({
            "stripePriceId": str(prid),
            "sku": r["SKU_INTERNO"],
            "variacao": r["VARIACAO_PRECO"],
            "valor": f"{float(r['VALOR_UNITARIO']):.2f}",
            "moeda": r["MOEDA"] or "BRL",
            "tipo": tipo_item(r["TIPO_PRECO_STRIPE"], tipo_receita),
            "recorrente": r["TIPO_PRECO_STRIPE"] == "Recorrente",
            "periodo": {"Mensal": "mensal", "Anual": "anual", "Único": "unico"}.get(r["PERIODO_FATURAMENTO_STRIPE"]),
            "revenueNature": natureza_receita(tipo_receita, cat),
            "tipoReceitaPlanilha": tipo_receita,
            "modeloCobranca": r["METADATA_modelo_cobranca"],
            "modeloPrecoStripe": r["MODELO_PRECO_STRIPE"],
            "exigeContrato": r["METADATA_exige_contrato"] == "sim",
            "descricaoPrice": r["DESCRICAO_DO_PRICE"],
            "lookupKey": r["CHAVE_DE_PESQUISA_PRICE"],
            "observacao": r["OBSERVACAO_OPERACIONAL"],
            "ordem": int(r["ORDEM"] or 0),
            "nomesLegado": nomes_legado(r["OBSERVACAO_OPERACIONAL"]),
        })

    # Ordena os prices de cada produto pela ORDEM da planilha (a sequência que
    # o Kauã montou), não por valor — "Starter, Medium, Large" é a ordem
    # comercial certa, ordenar por preço embaralharia isso.
    for p in produtos.values():
        p["precos"].sort(key=lambda x: x["ordem"])

    # Duplicidade de stripe_price_id ou lookup_key é ERRO: as duas são chave de
    # identidade do price (a primeira é a chave de idempotência do import).
    #
    # Duplicidade de SKU_INTERNO é AVISO, não erro. Caso real na planilha: dois
    # planos "PABX - Plano / Starter" legítimos com preços diferentes
    # (R$ 299,90 até 10 ramais e R$ 399,90 até 30 ramais, ambos herdados do
    # Zoho). O SKU colidiu, mas os lookup_key são distintos
    # (..._starter_299_mensal / ..._starter_399_mensal). Em vez de quebrar o
    # import ou editar a planilha do Kauã por conta própria, o SKU do segundo
    # em diante passa a ser o lookup_key (que já carrega o valor) e o caso é
    # reportado — a decisão de renomear no Stripe é comercial, não minha.
    vistos_price, vistos_sku, vistos_lookup = {}, {}, {}
    for p in produtos.values():
        for pr in p["precos"]:
            for campo, mapa in (("stripePriceId", vistos_price), ("lookupKey", vistos_lookup)):
                v = pr[campo]
                if not v:
                    problemas.append(f"{pr['stripePriceId']}: {campo} vazio")
                elif v in mapa:
                    problemas.append(f"{campo} duplicado: {v!r} em {mapa[v]} e {pr['stripePriceId']}")
                else:
                    mapa[v] = pr["stripePriceId"]

            sku = pr["sku"]
            if not sku:
                problemas.append(f"{pr['stripePriceId']}: SKU_INTERNO vazio")
            elif sku in vistos_sku:
                pr["skuOriginalPlanilha"] = sku
                pr["sku"] = pr["lookupKey"]
                avisos.append(
                    f"SKU '{sku}' repetido em {vistos_sku[sku]} e {pr['stripePriceId']} "
                    f"({p['nome']} / {pr['variacao']} R$ {pr['valor']}) — usando o lookup_key como SKU deste"
                )
            else:
                vistos_sku[sku] = pr["stripePriceId"]

    n_precos = sum(len(p["precos"]) for p in produtos.values())
    n_com_legado = sum(1 for p in produtos.values() for pr in p["precos"] if pr["nomesLegado"])
    saida = {
        "geradoDe": origem.name,
        "aba": "Catalogo_Unico",
        "totais": {"linhas": len(linhas), "produtos": len(produtos), "precos": n_precos, "precosComNomeLegado": n_com_legado},
        "ordemCategorias": ORDEM_CATEGORIA,
        "avisos": avisos,
        "produtos": list(produtos.values()),
    }
    DESTINO.parent.mkdir(parents=True, exist_ok=True)
    DESTINO.write_text(json.dumps(saida, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    print(f"planilha : {origem.name}")
    print(f"linhas   : {len(linhas)}")
    print(f"produtos : {len(produtos)}")
    print(f"prices   : {n_precos}")
    print(f"com nome legado do Zoho: {n_com_legado}")
    print(f"destino  : {DESTINO.relative_to(RAIZ)}")
    por_fam = collections.Counter(p["serviceFamily"] for p in produtos.values())
    print("por familia (produtos):", dict(por_fam))
    if avisos:
        print(f"\n{len(avisos)} AVISO(S) — importado, mas confira:")
        for x in avisos:
            print("  !", x)
    if problemas:
        print(f"\n{len(problemas)} PROBLEMA(S):")
        for x in problemas[:40]:
            print("  -", x)
        sys.exit(1)
    print("\nnenhum problema de integridade.")


if __name__ == "__main__":
    main()
